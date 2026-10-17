"""Demo project: merge monthly invoice CSVs into one summary Excel workbook.
Usage: python merge_invoices.py [input_folder] [output.xlsx]"""
import csv, glob, os, sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Reference

def main(src="sample_invoices", out="invoice_summary.xlsx"):
    rows=[]
    for p in sorted(glob.glob(os.path.join(src,"*.csv"))):
        with open(p,newline="") as f:
            for r in csv.DictReader(f):
                r["amount_gbp"]=float(r["amount_gbp"]); r["source_file"]=os.path.basename(p); rows.append(r)
    if not rows: sys.exit(f"No CSV files found in {src}")
    wb=Workbook(); ws=wb.active; ws.title="All invoices"
    hdr=["invoice_no","date","client","description","amount_gbp","status","source_file"]; ws.append(hdr)
    for r in sorted(rows,key=lambda r:r["date"]): ws.append([r[h] for h in hdr])
    for c in ws[1]: c.font=Font(bold=True)
    s=wb.create_sheet("By client"); s.append(["client","invoices","total_gbp","unpaid_gbp"])
    agg=defaultdict(lambda:[0,0.0,0.0])
    for r in rows:
        a=agg[r["client"]]; a[0]+=1; a[1]+=r["amount_gbp"]; a[2]+=r["amount_gbp"] if r["status"]=="unpaid" else 0
    for k,(n,t,u) in sorted(agg.items()): s.append([k,n,round(t,2),round(u,2)])
    n=len(agg)+1; s.append(["TOTAL",f"=SUM(B2:B{n})",f"=SUM(C2:C{n})",f"=SUM(D2:D{n})"])
    for c in s[1]+s[n+1]: c.font=Font(bold=True)
    ch=BarChart(); ch.title="Invoiced by client (GBP)"
    ch.add_data(Reference(s,min_col=3,min_row=1,max_row=n),titles_from_data=True); ch.set_categories(Reference(s,min_col=1,min_row=2,max_row=n))
    s.add_chart(ch,"F2"); wb.save(out)
    tot=sum(r["amount_gbp"] for r in rows); unp=sum(r["amount_gbp"] for r in rows if r["status"]=="unpaid")
    print(f"Merged {len(rows)} invoices from {len(set(r['source_file'] for r in rows))} files")
    for k,(n_,t,u) in sorted(agg.items()): print(f"  {k:<22} {n_:>2} invoices  £{t:>9,.2f}  unpaid £{u:>8,.2f}")
    print(f"Total £{tot:,.2f} | unpaid £{unp:,.2f} -> {out}")

if __name__=="__main__": main(*sys.argv[1:])
